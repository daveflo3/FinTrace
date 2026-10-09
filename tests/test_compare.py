from pathlib import Path

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

from fintrace.compare import ChangeClassification, ChangeKind, ModelComparator
from fintrace.excel import ExcelIngestor


def make_version(path: Path, *, growth: float, formula_multiplier: float = 0.75, add_output: bool = False) -> None:
    wb = Workbook()
    inputs = wb.active
    inputs.title = "Inputs"
    model = wb.create_sheet("Model")
    actuals = wb.create_sheet("Historical Actuals")

    inputs["A1"] = "Revenue"
    inputs["B1"] = 1000
    inputs["A2"] = "Revenue Growth"
    inputs["B2"] = growth
    inputs["B2"].number_format = "0.0%"
    wb.defined_names.add(DefinedName("RevenueGrowth", attr_text="'Inputs'!$B$2"))

    actuals["A1"] = "Reported Revenue"
    actuals["B1"] = 900 if growth < 0.07 else 925

    model["A1"] = "Forecast Revenue"
    model["B1"] = "='Inputs'!B1*(1+RevenueGrowth)"
    model["A2"] = "Free Cash Flow"
    model["B2"] = f"=B1*{formula_multiplier}"
    if add_output:
        model["A3"] = "Equity Value"
        model["B3"] = "=B2*10"

    wb.save(path)


def test_assumption_change_is_classified_and_traced(tmp_path: Path) -> None:
    old = tmp_path / "old.xlsx"
    new = tmp_path / "new.xlsx"
    make_version(old, growth=0.08)
    make_version(new, growth=0.06)

    report = ModelComparator().compare_files(old, new)
    change = next(item for item in report.changes if item.ref == "'Inputs'!B2")

    assert change.kind is ChangeKind.VALUE_CHANGED
    assert change.classification is ChangeClassification.ASSUMPTION
    assert "'Model'!B2" in change.impacted_outputs


def test_formula_change_is_detected(tmp_path: Path) -> None:
    old = tmp_path / "old.xlsx"
    new = tmp_path / "new.xlsx"
    make_version(old, growth=0.08, formula_multiplier=0.75)
    make_version(new, growth=0.08, formula_multiplier=0.70)

    report = ModelComparator().compare_files(old, new)
    change = next(item for item in report.changes if item.ref == "'Model'!B2")

    assert change.kind is ChangeKind.FORMULA_CHANGED
    assert change.classification is ChangeClassification.CALCULATION
    assert "0.75" in (change.before_formula or "")
    assert "0.7" in (change.after_formula or "")


def test_reported_actual_uses_transparent_heuristic(tmp_path: Path) -> None:
    old = tmp_path / "old.xlsx"
    new = tmp_path / "new.xlsx"
    make_version(old, growth=0.06)
    make_version(new, growth=0.08)

    report = ModelComparator().compare_files(old, new)
    change = next(item for item in report.changes if item.ref == "'Historical Actuals'!B1")

    assert change.classification is ChangeClassification.REPORTED_ACTUAL
    assert any("historical" in reason for reason in change.reasons)


def test_structural_addition_is_detected(tmp_path: Path) -> None:
    old = tmp_path / "old.xlsx"
    new = tmp_path / "new.xlsx"
    make_version(old, growth=0.08)
    make_version(new, growth=0.08, add_output=True)

    report = ModelComparator().compare_files(old, new)
    added = {item.ref: item for item in report.structural_changes if item.kind is ChangeKind.ADDED}

    assert "'Model'!B3" in added
    assert added["'Model'!B3"].classification is ChangeClassification.STRUCTURAL


def test_metadata_only_change_is_detected(tmp_path: Path) -> None:
    old = tmp_path / "old.xlsx"
    new = tmp_path / "new.xlsx"
    make_version(old, growth=0.08)
    make_version(new, growth=0.08)

    wb = __import__("openpyxl").load_workbook(new)
    wb["Inputs"]["B1"].number_format = "0.0"
    wb.save(new)

    report = ModelComparator().compare(
        ExcelIngestor().inspect(old),
        ExcelIngestor().inspect(new),
    )
    change = next(item for item in report.changes if item.ref == "'Inputs'!B1")
    assert change.kind is ChangeKind.METADATA_CHANGED
    assert change.classification is ChangeClassification.METADATA


def test_report_can_be_saved_as_json(tmp_path: Path) -> None:
    old = tmp_path / "old.xlsx"
    new = tmp_path / "new.xlsx"
    make_version(old, growth=0.08)
    make_version(new, growth=0.06)

    report = ModelComparator().compare_files(old, new)
    output = tmp_path / "comparison.json"
    report.save_json(output)

    payload = output.read_text(encoding="utf-8")
    assert '"candidate_assumption_change"' in payload
    assert '"old.xlsx"' in payload
    assert '"new.xlsx"' in payload
