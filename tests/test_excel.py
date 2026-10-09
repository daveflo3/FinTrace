from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

from fintrace.excel import ExcelIngestor, ExcelModelError


def make_workbook(path: Path) -> None:
    wb = Workbook()
    inputs = wb.active
    inputs.title = "Inputs"
    model = wb.create_sheet("Model")

    inputs["A1"] = "Revenue"
    inputs["B1"] = 1000
    inputs["A2"] = "Revenue Growth"
    inputs["B2"] = 0.08
    inputs["B2"].number_format = "0.0%"
    inputs["A3"] = "EBIT Margin"
    inputs["B3"] = 0.20
    inputs["B3"].number_format = "0.0%"

    wb.defined_names.add(DefinedName("RevenueGrowth", attr_text="'Inputs'!$B$2"))
    wb.defined_names.add(DefinedName("EBITMargin", attr_text="'Inputs'!$B$3"))

    model["A1"] = "Forecast Revenue"
    model["B1"] = "='Inputs'!B1*(1+RevenueGrowth)"
    model["A2"] = "EBIT"
    model["B2"] = "=B1*EBITMargin"
    model["A3"] = "Free Cash Flow"
    model["B3"] = "=B2*0.75"

    wb.save(path)


def test_inspection_maps_cross_sheet_and_named_range_dependencies(tmp_path: Path) -> None:
    path = tmp_path / "model.xlsx"
    make_workbook(path)

    snapshot = ExcelIngestor().inspect(path)
    edges = {(edge.upstream, edge.downstream) for edge in snapshot.dependencies}

    assert ("'Inputs'!B1", "'Model'!B1") in edges
    assert ("'Inputs'!B2", "'Model'!B1") in edges
    assert ("'Inputs'!B3", "'Model'!B2") in edges
    assert ("'Model'!B1", "'Model'!B2") in edges
    assert ("'Model'!B2", "'Model'!B3") in edges


def test_named_ranges_are_preserved(tmp_path: Path) -> None:
    path = tmp_path / "model.xlsx"
    make_workbook(path)

    snapshot = ExcelIngestor().inspect(path)
    names = {item.name: item.destinations for item in snapshot.named_ranges}

    assert names["RevenueGrowth"] == ("'Inputs'!B2",)
    assert "RevenueGrowth" in snapshot.cells["'Inputs'!B2"].named_ranges


def test_assumption_and_output_candidates_are_review_hints(tmp_path: Path) -> None:
    path = tmp_path / "model.xlsx"
    make_workbook(path)

    snapshot = ExcelIngestor().inspect(path)
    by_key = {(candidate.ref, candidate.role): candidate for candidate in snapshot.candidates}

    growth = by_key[("'Inputs'!B2", "assumption")]
    assert growth.score >= 0.8
    assert any("percentage" in reason for reason in growth.reasons)

    fcf = by_key[("'Model'!B3", "output")]
    assert fcf.score >= 0.8
    assert any("financial output" in reason for reason in fcf.reasons)


def test_snapshot_round_trip_json(tmp_path: Path) -> None:
    path = tmp_path / "model.xlsx"
    make_workbook(path)
    snapshot = ExcelIngestor().inspect(path)

    output = tmp_path / "map.json"
    snapshot.save_json(output)

    payload = output.read_text(encoding="utf-8")
    assert '"workbook_name": "model.xlsx"' in payload
    assert '"RevenueGrowth"' in payload


def test_large_ranges_are_not_silently_expanded(tmp_path: Path) -> None:
    path = tmp_path / "large.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Model"
    ws["A1"] = 1
    ws["B1"] = "=SUM(A1:A10)"
    wb.save(path)

    snapshot = ExcelIngestor(max_expanded_range_cells=5).inspect(path)
    assert snapshot.dependencies == ()
    assert any("exceeds the expansion limit" in warning for warning in snapshot.warnings)


def test_non_excel_input_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "model.csv"
    path.write_text("x,y\n1,2\n", encoding="utf-8")

    with pytest.raises(ExcelModelError, match="supports .xlsx and .xlsm"):
        ExcelIngestor().inspect(path)
